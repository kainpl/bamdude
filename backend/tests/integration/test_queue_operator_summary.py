"""Compact attention reads must not materialize maintenance or read job history."""

import pytest
from sqlalchemy import func, select

from backend.app.models.maintenance import MaintenanceType, PrinterMaintenance
from backend.app.models.printer_queue import PrinterQueue


@pytest.mark.asyncio
async def test_queue_attention_precedence_and_archived(async_client, printer_factory, db_session):
    printers = [await printer_factory(serial_number=f"SUMMARY{i}") for i in range(4)]
    for printer, status, paused in zip(
        printers, ["printing", "error", "idle", "error"], [True, True, False, False], strict=True
    ):
        db_session.add(PrinterQueue(id=printer.id, printer_id=printer.id, status=status, is_paused=paused))
    printers[-1].archived = True
    printers[0].is_active = False
    await db_session.commit()
    response = await async_client.get("/api/v1/queues/status-summary")
    assert response.status_code == 200
    data = response.json()
    assert data["attention"] == "error"
    assert {r["printer_id"]: r["effective_state"] for r in data["queues"]} == {
        printers[0].id: "paused",
        printers[1].id: "error",
        printers[2].id: "idle",
    }
    assert all("pending_count" not in row for row in data["queues"])


@pytest.mark.asyncio
async def test_maintenance_summary_is_read_only_and_matches_materialized_detail(
    async_client, printer_factory, db_session
):
    printer = await printer_factory(serial_number="SUMMARYMAINT", model="P1S")
    printer.runtime_seconds = 95 * 3600
    # A known system type exercises the virtual default without seeding catalogues.
    kind = MaintenanceType(
        name="Test rods",
        type_code="clean_carbon_rods",
        is_system=True,
        printer_models='["P1S"]',
        default_interval_hours=100,
    )
    db_session.add(kind)
    await db_session.commit()
    before_types = await db_session.scalar(select(func.count()).select_from(MaintenanceType))
    response = await async_client.get("/api/v1/maintenance/summary")
    assert response.status_code == 200
    assert response.json()["total_warning"] == 1
    assert await db_session.scalar(select(func.count()).select_from(PrinterMaintenance)) == 0
    assert await db_session.scalar(select(func.count()).select_from(MaintenanceType)) == before_types
    # Materialize exactly the same catalogue, without ensure_default_types adding unrelated defaults.
    db_session.add(
        PrinterMaintenance(
            printer_id=printer.id,
            maintenance_type_id=kind.id,
            enabled=True,
            last_performed_hours=0,
            custom_interval_hours=0,
        )
    )
    await db_session.commit()
    assert (await async_client.get("/api/v1/maintenance/summary")).json() == response.json()


@pytest.mark.asyncio
async def test_attention_and_history_do_not_leak_hidden_printers(async_client, printer_factory, db_session):
    import io

    from openpyxl import load_workbook

    from backend.app.models.maintenance import MaintenanceHistory
    from backend.app.models.user import User
    from backend.tests.integration.test_workshop_library_rights import _key

    visible = await printer_factory(serial_number="VISIBLESUMMARY", model="P1S")
    hidden = await printer_factory(serial_number="HIDDENSUMMARY", model="P1S")
    db_session.add_all(
        [
            PrinterQueue(id=visible.id, printer_id=visible.id),
            PrinterQueue(id=hidden.id, printer_id=hidden.id, status="error"),
        ]
    )
    kind = MaintenanceType(name="Private task", is_system=False, default_interval_hours=1)
    db_session.add(kind)
    await db_session.flush()
    item = PrinterMaintenance(printer_id=hidden.id, maintenance_type_id=kind.id, enabled=True, last_performed_hours=0)
    hidden.runtime_seconds = 10 * 3600
    db_session.add(item)
    await db_session.flush()
    db_session.add(MaintenanceHistory(printer_maintenance_id=item.id, hours_at_maintenance=10, notes="hidden note"))
    await db_session.commit()
    owner = await db_session.scalar(select(User).where(User.username == "test_admin"))
    raw = await _key(db_session, owner, can_read_status=True, printer_ids=[visible.id])
    headers = {"Authorization": f"Bearer {raw}"}
    response = await async_client.get("/api/v1/queues/status-summary", headers=headers)
    assert response.status_code == 200, response.text
    assert response.json()["attention"] == "none"
    assert [r["printer_id"] for r in response.json()["queues"]] == [visible.id]
    response = await async_client.get("/api/v1/maintenance/summary", headers=headers)
    assert response.json()["total_due"] == 0
    overview = await async_client.get("/api/v1/maintenance/overview", headers=headers)
    assert [row["printer_id"] for row in overview.json()] == [visible.id]
    history = await async_client.get("/api/v1/maintenance/history", headers=headers)
    assert history.json()["items"] == [] and history.json()["total"] == 0
    export = await async_client.get("/api/v1/maintenance/history/export", headers=headers)
    assert export.status_code == 200
    workbook = load_workbook(io.BytesIO(export.content))
    assert workbook.active.max_row == 1


@pytest.mark.asyncio
async def test_summary_query_count_does_not_scale_with_printers(async_client, printer_factory, db_session, test_engine):
    from sqlalchemy import event

    statements = []

    def record(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)

    await printer_factory(serial_number="QUERYCOUNT1")
    event.listen(test_engine.sync_engine, "before_cursor_execute", record)
    try:
        await async_client.get("/api/v1/maintenance/summary")
        first = sum(
            any(f"FROM {table}" in sql for table in ("printers", "printer_maintenance", "maintenance_types"))
            for sql in statements
        )
        for i in range(8):
            await printer_factory(serial_number=f"QUERYCOUNTMORE{i}")
        statements.clear()
        response = await async_client.get("/api/v1/maintenance/summary")
        assert response.status_code == 200
        second = sum(
            any(f"FROM {table}" in sql for table in ("printers", "printer_maintenance", "maintenance_types"))
            for sql in statements
        )
        assert first == second
        assert not any(sql.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")) for sql in statements)
    finally:
        event.remove(test_engine.sync_engine, "before_cursor_execute", record)


@pytest.mark.asyncio
@pytest.mark.parametrize("model,existing,expected", [("P1S", False, 0), ("P1S", True, 1), (None, False, 1)])
async def test_summary_detail_parity_for_calendar_defaults_and_model_exceptions(
    async_client, printer_factory, db_session, monkeypatch, model, existing, expected
):
    from unittest.mock import AsyncMock

    from backend.app.api.routes import maintenance

    monkeypatch.setattr(maintenance, "ensure_default_types", AsyncMock())
    printer = await printer_factory(serial_number="CALENDARPARITY", model=model)
    kind = MaintenanceType(
        name="Calendar task",
        type_code="test_calendar",
        is_system=True,
        default_interval_hours=90,
        interval_type="days",
        printer_models='["X1C"]',
    )
    db_session.add(kind)
    await db_session.flush()
    if existing:
        db_session.add(
            PrinterMaintenance(
                printer_id=printer.id,
                maintenance_type_id=kind.id,
                enabled=True,
                last_performed_hours=0,
                custom_interval_hours=0,
            )
        )
    await db_session.commit()
    summary = await async_client.get("/api/v1/maintenance/summary")
    assert summary.json()["total_due"] == expected
    detail = await async_client.get(f"/api/v1/maintenance/printers/{printer.id}")
    assert detail.json()["due_count"] == expected
    assert (await async_client.get("/api/v1/maintenance/summary")).json() == summary.json()
